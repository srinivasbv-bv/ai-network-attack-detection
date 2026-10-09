import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation_pptx():
    pptx_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "AI_Powered_Threat_Detection_Presentation.pptx")
    prs = Presentation()

    # Define color scheme
    NAVY = RGBColor(15, 23, 42)        # #0f172a
    CYAN = RGBColor(0, 242, 254)       # #00f2fe
    SLATE = RGBColor(30, 41, 59)       # #1e293b
    WHITE = RGBColor(255, 255, 255)
    GRAY = RGBColor(138, 153, 181)
    GREEN = RGBColor(0, 230, 118)
    ORANGE = RGBColor(255, 145, 0)

    # Set slide dimensions (16:9 widescreen)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    def add_slide_header(slide, title_text, category_text="AI POWERED THREAT DETECTION SYSTEM"):
        # Header Background
        header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = NAVY
        header_box.line.color.rgb = NAVY

        # Category
        tx_cat = slide.shapes.add_textbox(Inches(0.5), Inches(0.12), Inches(12), Inches(0.3))
        tf_cat = tx_cat.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = CYAN

        # Title
        tx_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12), Inches(0.6))
        tf_title = tx_title.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

    # SLIDE 1: Title Slide
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.color.rgb = NAVY

    tx_t1 = slide1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(2.0))
    p1 = tx_t1.text_frame.paragraphs[0]
    p1.text = "AI POWERED THREAT DETECTION SYSTEM"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = CYAN
    p1.alignment = PP_ALIGN.CENTER

    p1_sub = tx_t1.text_frame.add_paragraph()
    p1_sub.text = "Deep Neural Network (Deep MLP) & Random Forest Protocol Anomaly Pipeline"
    p1_sub.font.size = Pt(18)
    p1_sub.font.color.rgb = WHITE
    p1_sub.alignment = PP_ALIGN.CENTER

    tx_meta = slide1.shapes.add_textbox(Inches(1.0), Inches(5.0), Inches(11.333), Inches(1.5))
    p_m1 = tx_meta.text_frame.paragraphs[0]
    p_m1.text = "Master of Science in Computer Science / MCA Final Year Project"
    p_m1.font.size = Pt(14)
    p_m1.font.color.rgb = GRAY
    p_m1.alignment = PP_ALIGN.CENTER

    p_m2 = tx_meta.text_frame.add_paragraph()
    p_m2.text = "Karnataka State Open University (KSOU), Mysore"
    p_m2.font.size = Pt(13)
    p_m2.font.color.rgb = GRAY
    p_m2.alignment = PP_ALIGN.CENTER

    # SLIDE 2: Project Overview & Objectives
    slide2 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide2, "Project Overview & Core Objectives")

    tx2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5))
    tf2 = tx2.text_frame
    tf2.word_wrap = True

    items2 = [
        ("Real-Time Threat Detection:", " Detects DoS, DDoS, PortScan, BruteForce, and ARP Spoofing/MITM attacks instantly from live packet flows."),
        ("Dual-Engine AI Architecture:", " Combines a 3-layer Deep MLP Neural Network (128->64->32) for flow traffic and a Random Forest Classifier for ARP protocol anomalies."),
        ("Zero-Day Outlier Engine:", " Implements statistical distance heuristics to flag unclassified traffic patterns deviating >3x std dev from baseline."),
        ("Threat Framework Alignment:", " Maps every alert directly to MITRE ATT&CK techniques (T1499, T1498, T1046, T1110, T1557, T1204), Cyber Kill Chain, CAPEC, and OWASP Top 10."),
        ("Active Automated Firewall Mitigation:", " Automatically generates executable Linux IPTables, Fail2Ban, and ARPTables blocking scripts for immediate threat isolation."),
        ("ELK Stack Integration:", " Exporters normalize alerts into Elastic Common Schema (ECS 8.x) JSON for centralized Kibana analytics.")
    ]
    for head, body in items2:
        p = tf2.add_paragraph()
        p.space_after = Pt(14)
        run_h = p.add_run()
        run_h.text = "• " + head
        run_h.font.bold = True
        run_h.font.size = Pt(15)
        run_h.font.color.rgb = CYAN

        run_b = p.add_run()
        run_b.text = body
        run_b.font.size = Pt(14)
        run_b.font.color.rgb = SLATE

    # SLIDE 3: Problem Statement
    slide3 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide3, "Problem Statement: Limitations of Traditional Systems")

    tx3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5))
    tf3 = tx3.text_frame
    tf3.word_wrap = True

    items3 = [
        ("Inability to Detect Unknown Threats:", " Legacy signature-based IDS fail against obfuscated attack payloads and novel zero-day exploits."),
        ("High False Positive Rates:", " Traditional rule triggers generate massive false alarm fatigue, overwhelming SOC security analysts."),
        ("Lack of Mathematical Interpretability:", " Standard black-box deep learning models predict class labels without explaining specific feature metrics."),
        ("Protocol-Level Blind Spots:", " Flow-only systems miss stateless layer-2 protocol attacks like ARP cache poisoning."),
        ("Lack of Automated Mitigation Guidance:", " Alerts often notify analysts without providing exact, copy-pasteable firewall blocking commands.")
    ]
    for head, body in items3:
        p = tf3.add_paragraph()
        p.space_after = Pt(16)
        run_h = p.add_run()
        run_h.text = "• " + head
        run_h.font.bold = True
        run_h.font.size = Pt(15)
        run_h.font.color.rgb = ORANGE

        run_b = p.add_run()
        run_b.text = body
        run_b.font.size = Pt(14)
        run_b.font.color.rgb = SLATE

    # SLIDE 4: Benchmark Model Evaluation & Results
    slide4 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide4, "Benchmark Evaluation Results & Metrics")

    tx4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5))
    tf4 = tx4.text_frame
    tf4.word_wrap = True

    p4_intro = tf4.add_paragraph()
    p4_intro.text = "Measured performance on 14,000 benchmark samples (CICIDS2017 & UNSW-NB15 schema, 80/20 train/test split):"
    p4_intro.font.size = Pt(15)
    p4_intro.font.bold = True
    p4_intro.space_after = Pt(16)

    metrics_items = [
        ("Deep MLP Flow Model Accuracy:", " 99.75% | Precision: 99.75% | Recall: 99.75% | F1-Score: 99.75% | FPR: 0.06% | FNR: 0.25%"),
        ("Random Forest ARP Model Accuracy:", " 100.00% | Precision: 100.00% | Recall: 100.00% | F1-Score: 100.00% | FPR: 0.00% | FNR: 0.00%"),
        ("Zero-Day Outlier Engine Recall:", " 97.50% Recall on statistical feature anomalies exceeding 3x std dev from baseline."),
        ("Model Storage Footprint:", " Lightweight joblib binaries (~280 KB) optimized for low-latency serverless edge inference.")
    ]
    for head, body in metrics_items:
        p = tf4.add_paragraph()
        p.space_after = Pt(14)
        run_h = p.add_run()
        run_h.text = "• " + head
        run_h.font.bold = True
        run_h.font.size = Pt(15)
        run_h.font.color.rgb = GREEN

        run_b = p.add_run()
        run_b.text = body
        run_b.font.size = Pt(14)
        run_b.font.color.rgb = SLATE

    # SLIDE 5: Conclusion & Live Verification Link
    slide5 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide5, "Conclusion & Live Application Verification")

    tx5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5))
    tf5 = tx5.text_frame
    tf5.word_wrap = True

    p5_1 = tf5.add_paragraph()
    p5_1.text = "Key Project Achievements:"
    p5_1.font.size = Pt(18)
    p5_1.font.bold = True
    p5_1.font.color.rgb = CYAN
    p5_1.space_after = Pt(14)

    achievements = [
        "Delivered a production-ready, full-stack AI Threat Detection System with dual DL/ML model fusion.",
        "Integrated automated active Linux Firewall mitigation command generation (IPTables / Fail2Ban).",
        "Achieved 99.75% flow classification accuracy and 100.00% ARP anomaly precision.",
        "Deployed globally on Vercel Serverless with fixed left sidebar navigation and live stream simulator."
    ]
    for item in achievements:
        p = tf5.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(15)
        p.font.color.rgb = SLATE
        p.space_after = Pt(10)

    p5_url = tf5.add_paragraph()
    p5_url.space_before = Pt(20)
    run_u1 = p5_url.add_run()
    run_u1.text = "Live Official Project URL: "
    run_u1.font.bold = True
    run_u1.font.size = Pt(16)
    run_u1.font.color.rgb = NAVY

    run_u2 = p5_url.add_run()
    run_u2.text = "https://ai-powered-threat-detection-system.vercel.app"
    run_u2.font.bold = True
    run_u2.font.size = Pt(16)
    run_u2.font.color.rgb = CYAN

    prs.save(pptx_path)
    print(f"[PPTX Generator] Successfully generated presentation: {pptx_path}")

if __name__ == "__main__":
    create_presentation_pptx()
